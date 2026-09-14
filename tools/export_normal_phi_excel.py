#!/usr/bin/env python3
"""Export the verified heuristic-normal BCH subset to a PHI Excel workbook."""

import csv
import os
import re
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LABELS = Path("/temp_work/ch278233/BCH_DATASET/LABELS27_V2/labels.csv")
COMBINED_REPORTS = Path("/temp_work/ch278233/COMBINED_reports.csv")
VOLUME_MAP = Path("/temp_work/ch278233/BCH_DATASET/LABELS/volume_map.tsv")
RAW_PHI = Path("/temp_work/ch278233/BCH_DATASET/reports_8k.csv")
OUTPUT = Path("/home/ch278233/pediatric_normal_0_to_18_PHI.xlsx")


def read_dicts(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def parse_excel_date(value: str):
    """Convert an Excel serial date (or common date string) to a date value."""
    value = value.strip()
    try:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    except ValueError:
        pass
    for date_format in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, date_format).date()
        except ValueError:
            continue
    raise ValueError("Unrecognized exam-date format")


def main() -> None:
    labels = read_dicts(LABELS)
    label_columns = [column for column in labels[0] if column != "VolumeName"]
    zero_positive = {
        row["VolumeName"]
        for row in labels
        if not any(float(row[column] or 0) > 0 for column in label_columns)
    }
    assert len(labels) == 8_813
    assert len(zero_positive) == 727

    reports = {row["VolumeName"]: row for row in read_dicts(COMBINED_REPORTS)}
    normal_pattern = re.compile(r"\bnormal\b|\bno acute abnormalit", re.IGNORECASE)
    heuristic_normal = {
        volume
        for volume in zero_positive
        if normal_pattern.search(reports[volume].get("Impressions_EN", ""))
    }
    assert len(heuristic_normal) == 213

    volume_to_accession = {
        row["volume_name"]: row["accession"]
        for row in read_dicts(VOLUME_MAP, delimiter="\t")
    }
    raw_by_accession: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in read_dicts(RAW_PHI):
        raw_by_accession[row["Accession Number"].strip()].append(row)

    selected: list[tuple[float, str, str, str, str]] = []
    for volume in heuristic_normal:
        accession = volume_to_accession[volume].strip()
        matches = raw_by_accession[accession]
        assert len(matches) == 1, f"Expected one PHI row for mapped volume; got {len(matches)}"
        row = matches[0]
        age = float(row["Patient Age"].strip())
        if 0 <= age <= 18:
            selected.append(
                (
                    age,
                    parse_excel_date(row["Exam Started Date"]),
                    row["Patient MRN"].strip(),
                    row["Patient Sex"].strip(),
                    row["Report Text"],
                )
            )

    selected.sort(key=lambda item: item[0])
    assert len(selected) == 180

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Age 0-18"
    sheet.append(["Age", "Exam Date", "MRN", "Sex", "Report"])

    for age, exam_date, mrn, sex, report in selected:
        sheet.append([age, exam_date, mrn, sex, report])
        # Preserve identifiers and free text as literal strings, not Excel formulas.
        sheet.cell(sheet.max_row, 3).data_type = "s"
        sheet.cell(sheet.max_row, 4).data_type = "s"
        sheet.cell(sheet.max_row, 5).data_type = "s"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "0.00"
        sheet.cell(row, 2).number_format = "mm/dd/yyyy"
        sheet.cell(row, 3).number_format = "@"
        sheet.cell(row, 5).alignment = Alignment(wrap_text=True, vertical="top")

    widths = {1: 10, 2: 22, 3: 18, 4: 10, 5: 120}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"

    workbook.properties.title = "BCH heuristic-normal studies, age 0-18"
    workbook.properties.subject = "Contains protected health information (PHI)"
    workbook.save(OUTPUT)
    os.chmod(OUTPUT, 0o600)

    # Re-open and validate the persisted workbook without emitting PHI.
    check = load_workbook(OUTPUT, read_only=True, data_only=False)
    persisted = check["Age 0-18"]
    assert persisted.max_row == 181
    assert persisted.max_column == 5
    assert [cell.value for cell in next(persisted.iter_rows(max_row=1))] == [
        "Age",
        "Exam Date",
        "MRN",
        "Sex",
        "Report",
    ]
    ages = [row[0].value for row in persisted.iter_rows(min_row=2, max_col=1)]
    assert ages == sorted(ages)
    assert min(ages) >= 0 and max(ages) <= 18
    check.close()

    print(f"Created {OUTPUT} with {len(selected)} PHI rows (mode 0600).")


if __name__ == "__main__":
    main()
